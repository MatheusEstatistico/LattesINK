"""
Coleta de dados do Currículo Lattes, sem interface gráfica.

Tudo o que não depende do Tkinter fica aqui:
  - leitura e validação dos IDs (.txt, .csv, .xlsx);
  - o navegador (Selenium) e a extração de nome e data de atualização;
  - o carimbo de data/hora nos prints;
  - a gravação da planilha (.xlsx) e do backup linha a linha (.csv).

Assim dá para testar essas partes sem abrir a janela.
"""

from __future__ import annotations

import codecs
import csv
import io
import logging
import os
import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from selenium import webdriver
from selenium.common.exceptions import (
    InvalidSessionIdException,
    JavascriptException,
    NoSuchWindowException,
    WebDriverException,
)
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By


# ---------------------------------------------------------------------------
# Configurações
# ---------------------------------------------------------------------------

URL_LATTES = "https://lattes.cnpq.br/{id}"

# No layout atual do Lattes o nome do pesquisador fica em <h2 class="nome">.
# Esse elemento só existe na página do currículo (não na do CAPTCHA), então
# ele também serve para saber que o CAPTCHA foi resolvido.
SELETOR_NOME = "h2.nome"

# Elementos que indicam uma página de CAPTCHA: reCAPTCHA, hCaptcha, Turnstile
# (Cloudflare) e CAPTCHAs de imagem costumam ter "captcha" ou "data-sitekey"
# em algum atributo. O "i" no fim de cada seletor ignora maiúsculas/minúsculas.
SELETOR_CAPTCHA = ", ".join([
    "iframe[src*='captcha' i]",
    "iframe[src*='challenges.cloudflare.com' i]",
    "[data-sitekey]",
    "[class*='captcha' i]",
    "[id*='captcha' i]",
    "[name*='captcha' i]",
    "[class*='turnstile' i]",
])

TEMPO_MAX_CAPTCHA = 180          # segundos esperando o usuário resolver o CAPTCHA
TEMPO_MAX_CARREGAMENTO = 60      # segundos para uma página carregar
TEMPO_PAGINA_INESPERADA = 30     # segundos numa página sem currículo e sem CAPTCHA antes de desistir
LIMITE_FALHAS_SEGUIDAS = 3       # currículos seguidos sem abrir antes de pausar a captura
LIMITE_DIAS_DESATUALIZADO = 365  # acima disso a planilha destaca o currículo

NOME_PLANILHA = "dados_lattes.xlsx"
NOME_CSV = "dados_lattes.csv"
NOME_LOG = "log.txt"

# Um ID Lattes tem exatamente 16 dígitos. Os "lookarounds" impedem que um
# número maior (17+ dígitos) seja aceito pela metade.
PADRAO_ID = re.compile(r"(?<!\d)\d{16}(?!\d)")

# Texto exibido no topo do currículo: "Última atualização do currículo em 15/03/2024"
PADRAO_DATA = re.compile(
    r"[úu]ltima\s+atualiza[çc][ãa]o\s+do\s+curr[íi]culo\s+em\s*(\d{2}/\d{2}/\d{4})",
    re.IGNORECASE,
)

STATUS_OK = "OK"
STATUS_SEM_DATA = "Data não encontrada"
STATUS_CAPTCHA = "CAPTCHA não resolvido"
STATUS_PAGINA_INESPERADA = "Página inesperada"
STATUS_ERRO = "Erro"

# O que o navegador está mostrando enquanto o programa espera o currículo
PAGINA_CARREGANDO = "carregando"
PAGINA_CURRICULO = "curriculo"
PAGINA_CAPTCHA = "captcha"
PAGINA_OUTRA = "outra"


# Avisos e erros com o traceback completo. A janela grava isso em lattesink.log;
# sem configuração nenhuma, só os erros aparecem (no terminal).
logger = logging.getLogger("lattesink.coletor")


class CapturaCancelada(Exception):
    """O usuário pediu para cancelar."""


class ErroNavegador(Exception):
    """Não foi possível abrir o Chrome (sem ele, nenhum ID pode ser capturado)."""


# ---------------------------------------------------------------------------
# Registro de um pesquisador
# ---------------------------------------------------------------------------

@dataclass
class Registro:
    """Resultado da captura de um ID. Todo registro tem sempre os mesmos campos."""

    id_lattes: str
    status: str
    nome: str = ""
    ultima_atualizacao: date | None = None
    arquivo_print: str = ""      # caminho completo do .png ("" se não houve print)
    observacao: str = ""
    capturado_em: datetime = field(default_factory=datetime.now)

    @property
    def link(self) -> str:
        return URL_LATTES.format(id=self.id_lattes)

    @property
    def abriu(self) -> bool:
        """True se o currículo chegou a abrir (mesmo que a data não tenha sido encontrada)."""
        return self.status in (STATUS_OK, STATUS_SEM_DATA)

    @property
    def dias_desde_atualizacao(self) -> int | None:
        """Dias entre a última atualização e o momento da captura."""
        if self.ultima_atualizacao is None:
            return None
        return (self.capturado_em.date() - self.ultima_atualizacao).days


# ---------------------------------------------------------------------------
# Leitura e validação dos IDs
# ---------------------------------------------------------------------------

def normalizar_ids(valores):
    """
    Extrai os IDs Lattes de uma lista de valores.

    Aceita tanto o ID puro quanto um link ("http://lattes.cnpq.br/1234..."),
    e todos os IDs de um mesmo valor (ex.: uma linha "ID1, ID2, ID3").
    Retorna (ids, invalidos, qtd_duplicados):
      - ids: IDs únicos, na ordem em que apareceram;
      - invalidos: valores não vazios que não contêm um ID de 16 dígitos;
      - qtd_duplicados: quantos IDs repetidos foram descartados.
    """
    ids, invalidos = [], []
    for valor in valores:
        texto = str(valor).strip()
        if not texto:
            continue
        encontrados = PADRAO_ID.findall(texto)
        if encontrados:
            ids.extend(encontrados)
        else:
            invalidos.append(texto)
    unicos = list(dict.fromkeys(ids))  # remove duplicados mantendo a ordem
    return unicos, invalidos, len(ids) - len(unicos)


def ler_texto(caminho):
    """
    Lê um arquivo de texto descobrindo a codificação.

    Arquivos com BOM de UTF-16 ou UTF-32 (o "Texto Unicode" do Excel e o ">"
    do PowerShell gravam assim) são lidos nessa codificação. Os demais tentam
    UTF-8 (com ou sem BOM, o que resolve o caractere invisível que o Bloco de
    Notas coloca no início) e depois o padrão do Excel/Windows em português.
    """
    with open(caminho, "rb") as f:
        dados = f.read()

    # O BOM do UTF-32 LE começa com o do UTF-16 LE, então ele é testado antes
    if dados.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        tentativas = ("utf-32",)
    elif dados.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        tentativas = ("utf-16",)
    else:
        tentativas = ("utf-8-sig", "cp1252", "latin-1")

    for codificacao in tentativas:
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ValueError("Não foi possível identificar a codificação do arquivo.")


def ler_ids_txt(caminho):
    """Lê um .txt com um ID (ou link) por linha. Retorna a lista de linhas não vazias."""
    return [linha.strip() for linha in ler_texto(caminho).splitlines() if linha.strip()]


@dataclass
class Coluna:
    """Uma coluna de planilha/CSV, usada na janela de seleção de coluna."""

    nome: str
    valores: list
    # True quando a coluna tem números grandes, o que indica IDs salvos como
    # número no Excel. O Excel guarda só 15 dígitos significativos, então o
    # 16º dígito desses IDs pode ter virado zero.
    tem_numeros_grandes: bool = False


def _celula_para_texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _e_numero_grande(valor):
    return (isinstance(valor, (int, float)) and not isinstance(valor, bool)
            and abs(valor) >= 1e14)


def _montar_colunas(linhas):
    """Transforma uma lista de linhas em colunas, detectando se há cabeçalho."""
    linhas = [list(linha) for linha in linhas if any(v not in (None, "") for v in linha)]
    if not linhas:
        return []

    largura = max(len(linha) for linha in linhas)
    for linha in linhas:
        linha.extend([None] * (largura - len(linha)))

    # Se a primeira linha já contém um ID, o arquivo não tem cabeçalho e ela
    # é dado (antes, esse primeiro ID era descartado sem aviso).
    primeira = linhas[0]
    tem_cabecalho = not any(PADRAO_ID.search(_celula_para_texto(v)) for v in primeira)
    dados = linhas[1:] if tem_cabecalho else linhas

    colunas = []
    for i in range(largura):
        titulo = _celula_para_texto(primeira[i]) if tem_cabecalho else ""
        brutos = [linha[i] for linha in dados]
        valores = [_celula_para_texto(v) for v in brutos]
        if not titulo and not any(valores):
            continue  # coluna totalmente vazia
        colunas.append(Coluna(
            nome=titulo or f"Coluna {get_column_letter(i + 1)}",
            valores=valores,
            tem_numeros_grandes=any(_e_numero_grande(v) for v in brutos),
        ))
    return colunas


def ler_colunas(caminho):
    """
    Lê um .xlsx ou .csv e retorna a lista de colunas (objetos Coluna).

    No .xlsx são lidas todas as abas. Quando mais de uma aba tem dados, o nome
    de cada coluna leva o nome da aba na frente (ex.: "Docentes › ID Lattes").
    """
    extensao = os.path.splitext(caminho)[1].lower()

    if extensao == ".xlsx":
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        try:
            abas = [(ws.title, _montar_colunas(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
        finally:
            wb.close()
        abas = [(titulo, colunas) for titulo, colunas in abas if colunas]
        if len(abas) > 1:
            for titulo, colunas in abas:
                for coluna in colunas:
                    coluna.nome = f"{titulo} › {coluna.nome}"
        return [coluna for _, colunas in abas for coluna in colunas]

    if extensao == ".csv":
        texto = ler_texto(caminho)
        try:
            # CSV salvo pelo Excel em português usa ";" como separador
            delimitador = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delimitador = ","  # uma coluna só
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        return _montar_colunas(linhas)

    raise ValueError(f"Formato não suportado: {extensao}")


# ---------------------------------------------------------------------------
# Extração de dados da página
# ---------------------------------------------------------------------------

def extrair_data_atualizacao(texto_pagina):
    """Retorna a data de 'Última atualização do currículo em ...' ou None."""
    encontrado = PADRAO_DATA.search(texto_pagina)
    if not encontrado:
        return None
    try:
        return datetime.strptime(encontrado.group(1), "%d/%m/%Y").date()
    except ValueError:
        return None


def texto_da_pagina(driver):
    """
    Texto visível da página inteira.

    Usa document.body.innerText direto no navegador: em currículos longos é
    dezenas de vezes mais rápido que find_element(By.TAG_NAME, "body").text,
    que passa pelo protocolo do WebDriver.
    """
    return driver.execute_script("return document.body ? document.body.innerText : '';") or ""


def descrever_pagina(driver, limite=150):
    """Título e começo do texto da página, para explicar na planilha por que o currículo não abriu."""
    try:
        titulo = (driver.title or "").strip()
        texto = " ".join(texto_da_pagina(driver).split())
    except WebDriverException:
        return "não foi possível ler a página"
    descricao = " — ".join(parte for parte in (titulo, texto) if parte) or "página em branco"
    return descricao if len(descricao) <= limite else descricao[:limite - 1].rstrip() + "…"


def resumir_erro(erro):
    """Mensagem curta do erro (as do Selenium trazem stacktrace e links de documentação)."""
    if isinstance(erro, (InvalidSessionIdException, NoSuchWindowException)):
        return "O navegador foi fechado durante a captura"
    mensagem = getattr(erro, "msg", None) or str(erro)
    linhas = mensagem.strip().splitlines()
    resumo = linhas[0].split("; For documentation")[0].strip() if linhas else ""
    return resumo[:200] or erro.__class__.__name__


# ---------------------------------------------------------------------------
# Print da tela
# ---------------------------------------------------------------------------

def _carregar_fonte(tamanho):
    for nome in ("arialbd.ttf", "DejaVuSans-Bold.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
        try:
            return ImageFont.truetype(nome, tamanho)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=tamanho)  # Pillow 10.1+
    except TypeError:
        return ImageFont.load_default()


def carimbar_data(caminho_arquivo, quando):
    """Escreve 'Capturado em: ...' no canto inferior direito do print."""
    texto = quando.strftime("Capturado em: %d/%m/%Y às %H:%M:%S")
    margem, espaco = 20, 10

    with Image.open(caminho_arquivo) as original:
        img = original.convert("RGBA")
    camada = Image.new("RGBA", img.size, (255, 255, 255, 0))
    desenho = ImageDraw.Draw(camada)
    fonte = _carregar_fonte(25)

    esq, topo, dir_, base = desenho.textbbox((0, 0), texto, font=fonte)
    largura, altura = dir_ - esq, base - topo
    x2, y2 = img.width - margem, img.height - margem
    x1, y1 = x2 - largura - 2 * espaco, y2 - altura - 2 * espaco

    desenho.rectangle([x1, y1, x2, y2], fill=(0, 0, 0, 180))
    desenho.text((x1 + espaco - esq, y1 + espaco - topo), texto,
                 fill=(255, 255, 255, 255), font=fonte)
    Image.alpha_composite(img, camada).convert("RGB").save(caminho_arquivo)


# ---------------------------------------------------------------------------
# Navegador
# ---------------------------------------------------------------------------

class ColetorLattes:
    """
    Mantém um único Chrome aberto durante toda a captura e visita um
    currículo de cada vez. Pode ser usado a partir de uma thread separada:
    não mexe em nada do Tkinter e só se comunica pela função `log`.
    """

    def __init__(self, cancelado, log=print, tempo_max_captcha=TEMPO_MAX_CAPTCHA,
                 tempo_pagina_inesperada=TEMPO_PAGINA_INESPERADA):
        self.cancelado = cancelado      # threading.Event
        self.log = log
        self.tempo_max_captcha = tempo_max_captcha
        self.tempo_pagina_inesperada = tempo_pagina_inesperada
        self.driver = None
        # Vira True na primeira vez que a página do CAPTCHA é reconhecida pelo
        # SELETOR_CAPTCHA. Só a partir daí o coletor desiste antes do tempo de
        # uma página que não é nem o CAPTCHA nem o currículo (ID inexistente,
        # erro do site...). Se o CAPTCHA nunca for reconhecido (por exemplo, se
        # o CNPq trocar de serviço), a espera continua sendo o tempo todo, como
        # antes, para nunca desistir enquanto o usuário resolve o CAPTCHA.
        self.captcha_reconhecido = False

    # --- ciclo de vida do navegador -------------------------------------

    def _criar_driver(self):
        opcoes = Options()
        opcoes.add_argument("--window-size=1080,720")
        driver = webdriver.Chrome(options=opcoes)
        driver.set_page_load_timeout(TEMPO_MAX_CARREGAMENTO)
        return driver

    def _garantir_driver(self):
        """Devolve o navegador aberto, abrindo um novo se o usuário tiver fechado o anterior."""
        if self.driver is not None:
            try:
                _ = self.driver.current_window_handle  # falha se a aba ou o navegador foi fechado
                return self.driver
            except NoSuchWindowException:
                # A aba usada foi fechada, mas o navegador continua aberto em outra
                try:
                    abas = self.driver.window_handles
                    if abas:
                        self.driver.switch_to.window(abas[0])
                        return self.driver
                except WebDriverException:
                    pass
            except WebDriverException:
                pass
            self.log("🔄  O navegador foi fechado. Abrindo outro...")
            self.fechar()

        if self.cancelado.is_set():
            raise CapturaCancelada()
        try:
            self.driver = self._criar_driver()
        except WebDriverException as erro:
            raise ErroNavegador(resumir_erro(erro)) from erro
        return self.driver

    def fechar(self):
        """Fecha o navegador. Pode ser chamado de qualquer thread, quantas vezes for."""
        driver, self.driver = self.driver, None
        if driver is not None:
            try:
                driver.quit()
            except Exception:
                pass

    # --- captura --------------------------------------------------------

    @staticmethod
    def _estado_pagina(driver):
        """O que o navegador está mostrando: uma das constantes PAGINA_*."""
        try:
            if driver.execute_script("return document.readyState") != "complete":
                return PAGINA_CARREGANDO
            if driver.find_elements(By.CSS_SELECTOR, SELETOR_NOME):
                return PAGINA_CURRICULO
            if driver.find_elements(By.CSS_SELECTOR, SELETOR_CAPTCHA):
                return PAGINA_CAPTCHA
        except JavascriptException:
            # A página trocou no meio da verificação (ex.: logo depois de resolver o CAPTCHA)
            return PAGINA_CARREGANDO
        return PAGINA_OUTRA

    def _aguardar_curriculo(self, driver):
        """
        Espera o currículo aparecer (ou seja, o CAPTCHA ser resolvido).

        Retorna o estado da página quando a espera termina: PAGINA_CURRICULO se
        o currículo abriu; senão, o que estava na tela quando o tempo acabou.
        Uma página que não é nem o currículo nem o CAPTCHA é abandonada depois
        de tempo_pagina_inesperada segundos (veja captcha_reconhecido).
        Levanta CapturaCancelada se o usuário cancelar.
        """
        fim = time.monotonic() + self.tempo_max_captcha
        outra_pagina_desde = None
        estado = PAGINA_CARREGANDO
        while time.monotonic() < fim:
            if self.cancelado.is_set():
                raise CapturaCancelada()
            estado = self._estado_pagina(driver)
            if estado == PAGINA_CURRICULO:
                return estado
            if estado == PAGINA_CAPTCHA:
                self.captcha_reconhecido = True
            if estado != PAGINA_OUTRA:
                outra_pagina_desde = None
            elif outra_pagina_desde is None:
                outra_pagina_desde = time.monotonic()
            elif (self.captcha_reconhecido
                    and time.monotonic() - outra_pagina_desde >= self.tempo_pagina_inesperada):
                return estado
            self.cancelado.wait(0.5)  # como time.sleep, mas acorda na hora se cancelar
        return estado

    def _registro_sem_curriculo(self, lattes_id, driver, estado):
        """Registro de um ID cujo currículo não abriu, com o motivo mais provável."""
        if estado == PAGINA_CARREGANDO:
            return Registro(lattes_id, STATUS_ERRO,
                            observacao=f"A página não terminou de carregar em {self.tempo_max_captcha} s")
        descricao = descrever_pagina(driver)
        if estado == PAGINA_OUTRA and self.captcha_reconhecido:
            return Registro(lattes_id, STATUS_PAGINA_INESPERADA,
                            observacao=f"A página não mostrou o currículo nem o CAPTCHA: {descricao}")
        observacao = f"O currículo não abriu em {self.tempo_max_captcha} s"
        if estado == PAGINA_OUTRA:
            # O CAPTCHA ainda não foi reconhecido nesta captura: pode ser ele ou outra página
            observacao += f" (página: {descricao})"
        return Registro(lattes_id, STATUS_CAPTCHA, observacao=observacao)

    def capturar(self, lattes_id, pasta):
        """
        Abre o currículo, espera o CAPTCHA, extrai os dados e salva o print.

        Sempre devolve um Registro (com status de erro, se for o caso).
        Só levanta exceção em dois casos: CapturaCancelada e ErroNavegador.
        """
        try:
            driver = self._garantir_driver()
            driver.get(URL_LATTES.format(id=lattes_id))
            limite = (f"{self.tempo_max_captcha // 60} min" if self.tempo_max_captcha >= 60
                      else f"{self.tempo_max_captcha} s")
            self.log(f"⏳  [{lattes_id}]  Aguardando o currículo abrir "
                     f"(resolva o CAPTCHA no navegador, até {limite})...")

            estado = self._aguardar_curriculo(driver)
            if estado != PAGINA_CURRICULO:
                return self._registro_sem_curriculo(lattes_id, driver, estado)

            elementos_nome = driver.find_elements(By.CSS_SELECTOR, SELETOR_NOME)
            nome = elementos_nome[0].text.strip() if elementos_nome else ""
            data = extrair_data_atualizacao(texto_da_pagina(driver))

            caminho_print = os.path.join(pasta, f"{lattes_id}.png")
            capturado_em = datetime.now()
            driver.save_screenshot(caminho_print)

            observacao = ""
            try:
                carimbar_data(caminho_print, capturado_em)
            except Exception as erro:
                observacao = f"Print sem data/hora: {resumir_erro(erro)}"
                self.log(f"⚠️  [{lattes_id}]  {observacao}")
                logger.warning("Não foi possível carimbar o print de %s", lattes_id, exc_info=True)

            return Registro(
                id_lattes=lattes_id,
                status=STATUS_OK if data else STATUS_SEM_DATA,
                nome=nome or "Não encontrado",
                ultima_atualizacao=data,
                arquivo_print=caminho_print,
                observacao=observacao,
                capturado_em=capturado_em,
            )

        except (CapturaCancelada, ErroNavegador):
            raise
        except Exception as erro:
            if self.cancelado.is_set():
                # O erro foi causado pelo próprio cancelamento (navegador fechado no meio)
                raise CapturaCancelada() from erro
            logger.warning("Erro ao capturar o currículo %s", lattes_id, exc_info=True)
            return Registro(lattes_id, STATUS_ERRO, observacao=resumir_erro(erro))


# ---------------------------------------------------------------------------
# Arquivos de saída
# ---------------------------------------------------------------------------

def criar_pasta_execucao(pasta_base, agora=None):
    """Cria uma subpasta nova para esta captura, ex.: captura_2026-09-22_1555."""
    agora = agora or datetime.now()
    nome = f"captura_{agora:%Y-%m-%d_%H%M}"
    caminho = os.path.join(pasta_base, nome)
    sufixo = 2
    while os.path.exists(caminho):
        caminho = os.path.join(pasta_base, f"{nome}_{sufixo}")
        sufixo += 1
    os.makedirs(caminho)
    return caminho


CABECALHOS_CSV = ["Link Lattes", "ID Lattes", "Nome do Pesquisador", "Última Atualização",
                  "Dias desde a atualização", "Status", "Observação", "Arquivo do print",
                  "Capturado em"]


def acrescentar_csv(registro, caminho_csv):
    """
    Acrescenta uma linha ao CSV de backup. O arquivo é aberto e fechado a cada
    registro, então nada se perde se o programa travar no meio da captura.
    Usa ";" e UTF-8 com BOM para abrir direto no Excel em português.
    """
    novo = not os.path.exists(caminho_csv)
    with open(caminho_csv, "a", newline="", encoding="utf-8-sig") as f:
        escritor = csv.writer(f, delimiter=";")
        if novo:
            escritor.writerow(CABECALHOS_CSV)
        data = registro.ultima_atualizacao
        dias = registro.dias_desde_atualizacao
        escritor.writerow([
            registro.link,
            registro.id_lattes,
            registro.nome,
            data.strftime("%d/%m/%Y") if data else "",
            "" if dias is None else dias,
            registro.status,
            registro.observacao,
            os.path.basename(registro.arquivo_print),
            registro.capturado_em.strftime("%d/%m/%Y %H:%M:%S"),
        ])


def _link_para_arquivo(caminho_arquivo, pasta_planilha):
    """Link relativo (continua funcionando se a pasta for movida), ou absoluto se não der."""
    try:
        return os.path.relpath(caminho_arquivo, pasta_planilha)
    except ValueError:  # Windows: arquivos em unidades diferentes (C: e D:)
        return os.path.abspath(caminho_arquivo)


def salvar_planilha(registros, caminho_xlsx):
    """Grava os registros em uma planilha Excel formatada."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Currículos Lattes"

    colunas = [  # (cabeçalho, largura)
        ("Link Lattes", 38),
        ("ID Lattes", 19),
        ("Nome do Pesquisador", 40),
        ("Última Atualização", 18),
        ("Dias desde a atualização", 15),
        ("Status", 22),
        ("Observação", 40),
        ("Arquivo do print", 22),
        ("Capturado em", 18),
    ]

    fonte_cabecalho = Font(bold=True, color="FFFFFF")
    fundo_cabecalho = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
    fonte_link = Font(color="0563C1", underline="single")
    fundo_problema = PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    esquerda = Alignment(horizontal="left", vertical="center")

    for col, (titulo, largura) in enumerate(colunas, start=1):
        celula = ws.cell(row=1, column=col, value=titulo)
        celula.font = fonte_cabecalho
        celula.fill = fundo_cabecalho
        celula.alignment = centro
        ws.column_dimensions[get_column_letter(col)].width = largura

    pasta_planilha = os.path.dirname(os.path.abspath(caminho_xlsx))

    for linha, r in enumerate(registros, start=2):
        valores = [
            r.link,
            r.id_lattes,               # texto, para o Excel não cortar dígitos
            r.nome,
            r.ultima_atualizacao,      # data de verdade (ordena e filtra)
            r.dias_desde_atualizacao,
            r.status,
            r.observacao,
            os.path.basename(r.arquivo_print),
            r.capturado_em,
        ]
        for col, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=col, value=valor)
            celula.alignment = esquerda if col in (3, 7) else centro

        ws.cell(row=linha, column=1).hyperlink = r.link
        ws.cell(row=linha, column=1).font = fonte_link
        ws.cell(row=linha, column=4).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=9).number_format = "DD/MM/YYYY HH:MM"

        if r.arquivo_print:
            ws.cell(row=linha, column=8).hyperlink = _link_para_arquivo(r.arquivo_print,
                                                                        pasta_planilha)
            ws.cell(row=linha, column=8).font = fonte_link

        if r.status != STATUS_OK:
            ws.cell(row=linha, column=6).fill = fundo_problema

    ultima_linha = max(len(registros) + 1, 2)

    # Destaca em vermelho os currículos sem atualização há mais de um ano
    ws.conditional_formatting.add(
        f"D2:E{ultima_linha}",
        FormulaRule(
            formula=[f"AND(ISNUMBER($E2),$E2>{LIMITE_DIAS_DESATUALIZADO})"],
            fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
            font=Font(color="9C0006", bold=True),
        ),
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{ultima_linha}"

    # Ao imprimir: paisagem, todas as colunas em uma página de largura
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    wb.save(caminho_xlsx)


# ---------------------------------------------------------------------------
# Captura em lote
# ---------------------------------------------------------------------------

def executar_lote(coletor_, pendentes, pasta, enviar, inicio=1, total=None):
    """
    Captura os IDs `pendentes`, um de cada vez, gravando o CSV de backup a cada um.

    Feito para rodar numa thread separada: não mexe na janela e só se comunica
    pela função enviar(tipo, dado), chamada com
      ("log", texto), ("registro", Registro) e, sempre no final,
      ("fim", (motivo, detalhe)), com motivo "concluido", "cancelado",
      "pausado" ou "erro".
    `inicio` e `total` só servem para numerar as mensagens ([inicio/total]).

    Depois de LIMITE_FALHAS_SEGUIDAS currículos seguidos sem abrir, a captura
    é pausada: provavelmente o site está fora do ar ou ninguém está resolvendo
    os CAPTCHAs, e continuar só faria o programa esperar à toa.
    """
    total = total or len(pendentes)
    caminho_csv = os.path.join(pasta, NOME_CSV)
    motivo, detalhe = "concluido", ""
    falhas_seguidas = 0
    try:
        for indice, lattes_id in enumerate(pendentes):
            posicao = inicio + indice
            if coletor_.cancelado.is_set():
                motivo = "cancelado"
                break
            enviar("log", f"🌐  [{posicao}/{total}]  Abrindo o currículo {lattes_id}")
            registro = coletor_.capturar(lattes_id, pasta)
            try:
                acrescentar_csv(registro, caminho_csv)
            except OSError as erro:
                enviar("log", "⚠️  Não foi possível gravar o CSV de backup "
                              f"(ele está aberto no Excel?): {erro}")
            enviar("registro", registro)

            falhas_seguidas = 0 if registro.abriu else falhas_seguidas + 1
            if falhas_seguidas >= LIMITE_FALHAS_SEGUIDAS and indice < len(pendentes) - 1:
                motivo = "pausado"
                detalhe = (f"{falhas_seguidas} currículos seguidos não abriram "
                           f"(o último: \"{registro.status}\").\n\n"
                           "O site do Lattes pode estar fora do ar ou ter mudado, ou os CAPTCHAs "
                           "não estão sendo resolvidos. A captura foi pausada para o programa "
                           "não ficar esperando à toa.")
                break
    except CapturaCancelada:
        motivo = "cancelado"
    except ErroNavegador as erro:
        logger.error("Não foi possível abrir o Chrome", exc_info=True)
        motivo, detalhe = "erro", (f"Não foi possível abrir o Chrome:\n{erro}\n\n"
                                   "Verifique se o Google Chrome está instalado e atualizado.")
    except Exception as erro:
        logger.exception("Erro inesperado durante a captura")
        motivo, detalhe = "erro", f"Erro inesperado: {resumir_erro(erro)}"
    finally:
        coletor_.fechar()
        enviar("fim", (motivo, detalhe))


# ---------------------------------------------------------------------------
# Autoteste do executável
# ---------------------------------------------------------------------------

def autoteste(pasta):
    """
    Confere, sem abrir o Chrome, se o que a captura usa está funcionando:
    carimbo no print, planilha, CSV, leitura de planilha e o selenium-manager.
    Feito para o executável gerado pelo PyInstaller (veja LattesINK.py --autoteste),
    onde falta de arquivo ou de módulo só aparece na hora de usar.
    Levanta uma exceção se algo falhar.
    """
    agora = datetime.now()
    caminho_print = os.path.join(pasta, "autoteste.png")
    Image.new("RGB", (400, 300), "white").save(caminho_print)
    carimbar_data(caminho_print, agora)

    registro = Registro("0000000000000000", STATUS_OK, nome="Autoteste", ultima_atualizacao=agora.date(),
                        arquivo_print=caminho_print, capturado_em=agora)
    planilha = os.path.join(pasta, NOME_PLANILHA)
    salvar_planilha([registro], planilha)
    acrescentar_csv(registro, os.path.join(pasta, NOME_CSV))
    lidos = [normalizar_ids(coluna.valores)[0] for coluna in ler_colunas(planilha)]
    if [registro.id_lattes] not in lidos:
        raise RuntimeError("a planilha gerada não pôde ser lida de volta")

    # O build deixa só o selenium-manager da própria plataforma: ele precisa estar lá
    binario = encontrar_selenium_manager()
    if not os.path.isfile(binario):
        raise RuntimeError(f"selenium-manager não encontrado em {binario}")


def encontrar_selenium_manager():
    """
    Caminho do selenium-manager que o Selenium vai usar nesta plataforma.

    A função do Selenium que faz isso é interna e já mudou de nome entre
    versões; se nenhum dos nomes conhecidos existir, procura o arquivo na
    pasta do próprio Selenium. Levanta uma exceção se não encontrar.
    """
    from selenium.webdriver.common import selenium_manager

    gerente = selenium_manager.SeleniumManager()
    for nome in ("_get_binary", "get_binary"):
        obter = getattr(gerente, nome, None)
        if callable(obter):
            return str(obter())

    for raiz, _, arquivos in os.walk(os.path.dirname(selenium_manager.__file__)):
        for arquivo in arquivos:
            if arquivo in ("selenium-manager", "selenium-manager.exe"):
                return os.path.join(raiz, arquivo)
    raise RuntimeError("selenium-manager não encontrado na pasta do Selenium")
