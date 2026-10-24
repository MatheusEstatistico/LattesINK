"""Arquivos gerados pela captura: subpasta, CSV de backup, planilha e carimbo no print."""

import csv
import os
from datetime import date, datetime

import openpyxl
from falsos import ID_A, ID_B
from PIL import Image

import coletor


CAPTURA = datetime(2026, 10, 16, 21, 30, 5)


def _registro_ok(pasta, **extra):
    dados = {
        "id_lattes": ID_A,
        "status": coletor.STATUS_OK,
        "nome": "Maria da Silva",
        "ultima_atualizacao": date(2024, 3, 15),
        "arquivo_print": os.path.join(pasta, f"{ID_A}.png"),
        "capturado_em": CAPTURA,
    }
    dados.update(extra)
    return coletor.Registro(**dados)


def _registro_erro():
    return coletor.Registro(ID_B, coletor.STATUS_ERRO, observacao="timeout", capturado_em=CAPTURA)


# --- Registro -----------------------------------------------------------------

def test_registro_calcula_link_e_dias_desde_a_atualizacao(tmp_path):
    r = _registro_ok(str(tmp_path))
    assert r.link == f"https://lattes.cnpq.br/{ID_A}"
    assert r.dias_desde_atualizacao == (CAPTURA.date() - date(2024, 3, 15)).days
    assert _registro_erro().dias_desde_atualizacao is None


# --- subpasta da captura --------------------------------------------------------

def test_cria_subpasta_sem_sobrescrever_a_anterior(tmp_path):
    primeira = coletor.criar_pasta_execucao(str(tmp_path), CAPTURA)
    segunda = coletor.criar_pasta_execucao(str(tmp_path), CAPTURA)
    assert os.path.basename(primeira) == "captura_2026-10-16_2130"
    assert os.path.basename(segunda) == "captura_2026-10-16_2130_2"
    assert os.path.isdir(primeira) and os.path.isdir(segunda)


# --- CSV de backup --------------------------------------------------------------

def test_csv_de_backup_abre_no_excel_em_portugues(tmp_path):
    caminho = tmp_path / "dados_lattes.csv"
    coletor.acrescentar_csv(_registro_ok(str(tmp_path)), str(caminho))
    coletor.acrescentar_csv(_registro_erro(), str(caminho))

    bruto = caminho.read_bytes()
    assert bruto.startswith(b"\xef\xbb\xbf")             # BOM: o Excel reconhece o UTF-8
    assert bruto.count(b"\xef\xbb\xbf") == 1             # cabeçalho e BOM só uma vez

    with open(caminho, encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    assert linhas[0] == coletor.CABECALHOS_CSV
    dias = str((CAPTURA.date() - date(2024, 3, 15)).days)
    assert linhas[1][1:6] == [ID_A, "Maria da Silva", "15/03/2024", dias, coletor.STATUS_OK]
    assert linhas[1][7:] == [f"{ID_A}.png", "16/10/2026 21:30:05"]
    assert linhas[2][1:7] == [ID_B, "", "", "", coletor.STATUS_ERRO, "timeout"]


# --- planilha -----------------------------------------------------------------

def test_planilha_formatada(tmp_path):
    caminho = tmp_path / "dados_lattes.xlsx"
    antigo = _registro_ok(str(tmp_path), id_lattes="1111222233334444", arquivo_print="",
                          ultima_atualizacao=date(2020, 1, 1))
    coletor.salvar_planilha([_registro_ok(str(tmp_path)), _registro_erro(), antigo], str(caminho))

    ws = openpyxl.load_workbook(caminho).active
    assert [c.value for c in ws[1]] == [
        "Link Lattes", "ID Lattes", "Nome do Pesquisador", "Última Atualização",
        "Dias desde a atualização", "Status", "Observação", "Arquivo do print", "Capturado em",
    ]
    assert ws.max_row == 4

    # ID como texto (o Excel não corta dígitos) e data como data de verdade
    assert ws["B2"].value == ID_A and ws["B2"].data_type == "s"
    assert ws["D2"].value.date() == date(2024, 3, 15)
    assert ws["D2"].number_format == "DD/MM/YYYY"

    # links para o currículo e para o print (relativo à pasta da planilha)
    assert ws["A2"].hyperlink.target == f"https://lattes.cnpq.br/{ID_A}"
    assert ws["H2"].hyperlink.target == f"{ID_A}.png"
    assert ws["H4"].hyperlink is None

    # status com problema destacado; currículos desatualizados com formatação condicional
    assert ws["F2"].fill.fgColor.rgb != ws["F3"].fill.fgColor.rgb
    regras = {str(cf.sqref): [r.formula for r in cf.rules] for cf in ws.conditional_formatting}
    assert regras == {"D2:E4": [[f"AND(ISNUMBER($E2),$E2>{coletor.LIMITE_DIAS_DESATUALIZADO})"]]}

    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:I4"


def test_planilha_sem_registros(tmp_path):
    caminho = tmp_path / "vazia.xlsx"
    coletor.salvar_planilha([], str(caminho))
    assert openpyxl.load_workbook(caminho).active.max_row == 1


# --- carimbo de data/hora no print ------------------------------------------------

def test_carimbo_no_canto_inferior_direito(tmp_path):
    caminho = tmp_path / "print.png"
    Image.new("RGB", (1080, 720), "white").save(caminho)

    coletor.carimbar_data(str(caminho), CAPTURA)

    with Image.open(caminho) as img:
        assert img.size == (1080, 720)
        assert img.getpixel((10, 10)) == (255, 255, 255)        # resto do print intacto
        assert img.getpixel((1080 - 25, 720 - 25)) != (255, 255, 255)  # caixa escura do carimbo


# --- autoteste do executável --------------------------------------------------------

def test_autoteste_gera_print_planilha_e_csv(tmp_path):
    coletor.autoteste(str(tmp_path))
    assert {"autoteste.png", coletor.NOME_PLANILHA, coletor.NOME_CSV} <= set(os.listdir(tmp_path))


def test_encontra_o_selenium_manager_mesmo_se_o_selenium_mudar_a_funcao(monkeypatch):
    from selenium.webdriver.common import selenium_manager

    caminho = coletor.encontrar_selenium_manager()
    assert os.path.isfile(caminho)

    class SemAFuncao:  # simula uma versão futura do Selenium, sem _get_binary/get_binary
        pass

    monkeypatch.setattr(selenium_manager, "SeleniumManager", SemAFuncao)
    alternativo = coletor.encontrar_selenium_manager()
    assert os.path.isfile(alternativo)
    assert os.path.basename(alternativo) in ("selenium-manager", "selenium-manager.exe")
