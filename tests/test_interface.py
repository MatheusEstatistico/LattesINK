"""
Testes da janela (Tkinter) com um coletor falso: nenhum teste abre o Chrome.

Precisam de um display. No Linux sem interface gráfica, rode com
"xvfb-run -a python -m pytest"; sem display, estes testes são pulados.
"""

import os
import time
from typing import ClassVar

import pytest


tk = pytest.importorskip("tkinter")

import openpyxl  # noqa: E402
from falsos import ID_A, ID_B, ID_C, ColetorFalso  # noqa: E402

import coletor  # noqa: E402
import LattesINK  # noqa: E402


@pytest.fixture
def avisos(monkeypatch):
    """Troca as caixas de diálogo por uma lista (senão os testes travariam esperando um clique)."""
    lista = []
    for funcao in ("showinfo", "showwarning", "showerror"):
        monkeypatch.setattr(LattesINK.messagebox, funcao,
                            lambda titulo, texto, _f=funcao, **_: lista.append((_f, titulo, texto)))
    for funcao in ("askokcancel", "askyesno"):
        monkeypatch.setattr(LattesINK.messagebox, funcao, lambda *_, **__: True)
    return lista


@pytest.fixture
def coletor_falso(monkeypatch):
    class Coletor(ColetorFalso):  # cada teste ganha o seu roteiro
        roteiro: ClassVar[dict] = {}
        bloquear: ClassVar[bool] = False

    monkeypatch.setattr(coletor, "ColetorLattes", Coletor)
    return Coletor


@pytest.fixture
def abrir_janela(monkeypatch, tmp_path, avisos, coletor_falso):
    """Abre janelas com a configuração isolada numa pasta temporária e as fecha no fim do teste."""
    monkeypatch.setattr(LattesINK, "ARQUIVO_CONFIG", str(tmp_path / "config" / "config.json"))
    monkeypatch.setattr(LattesINK, "ARQUIVO_CONFIG_ANTIGO", str(tmp_path / ".lattesink.json"))
    abertas = []

    def abrir():
        try:
            janela = LattesINK.LattesInk()
        except tk.TclError as erro:
            pytest.skip(f"sem display para abrir a janela: {erro}")
        abertas.append(janela)
        janela.janela.update()
        return janela

    yield abrir
    for janela in abertas:
        janela.janela.after_cancel(janela._id_fila)
        janela.janela.destroy()


@pytest.fixture
def app(abrir_janela):
    return abrir_janela()


def preparar(app, pasta, ids=(ID_A, ID_B, ID_C)):
    os.makedirs(pasta, exist_ok=True)
    app._definir_pasta(str(pasta), lembrar=False)
    assert app._aplicar_ids(list(ids), "ids.txt")


def esperar_fim(app, limite=20):
    fim = time.monotonic() + limite
    while app.captura_ativa and time.monotonic() < fim:
        app.janela.update()
        time.sleep(0.02)
    app.janela.update()
    assert not app.captura_ativa, "a captura não terminou a tempo"


def ativo(botao):
    return str(botao.cget("state")) == "normal"


def test_titulo_com_a_versao(app):
    assert app.janela.title() == f"LattesINK {LattesINK.__version__}"
    assert app.btn_ids.cget("text") == "Lattes IDs"


def test_janela_comporta_todo_o_conteudo(app):
    j = app.janela
    j.update()
    assert j.winfo_width() >= j.winfo_reqwidth()
    assert j.winfo_height() >= j.winfo_reqheight()
    assert j.minsize() == (j.winfo_reqwidth(), j.winfo_reqheight())
    assert j.resizable() == (True, True)


def test_medidas_em_pixels_acompanham_a_escala_da_tela(monkeypatch, abrir_janela):
    # Simula uma tela com escala de 150% (144 pontos por polegada)
    monkeypatch.setattr(tk.Tk, "winfo_fpixels", lambda self, distancia: 144.0)
    app = abrir_janela()
    assert app.escala == 1.5
    assert app._px(560) == 840
    assert int(app.label_pasta.cget("wraplength")) == 840


def test_estado_inicial(app):
    assert not ativo(app.btn_iniciar)
    assert not ativo(app.btn_cancelar)
    assert not ativo(app.btn_exportar)
    assert ativo(app.btn_pasta) and ativo(app.btn_ids)


def test_botao_desativado_fica_cinza(app, tmp_path):
    assert app.btn_iniciar.cget("bg") == LattesINK.COR_DESATIVADO
    preparar(app, tmp_path / "saida")
    assert app.btn_iniciar.cget("bg") == LattesINK.COR_INICIAR  # volta à cor original
    assert app.btn_cancelar.cget("bg") == LattesINK.COR_DESATIVADO


def test_ids_carregados_habilitam_a_captura(app, tmp_path):
    preparar(app, tmp_path / "saida")
    assert ativo(app.btn_iniciar)
    assert app.label_progresso.cget("text") == "0 / 3"
    assert "3 ID(s) carregado(s)" in app.label_ids.cget("text")


def test_captura_completa(app, tmp_path, avisos):
    preparar(app, tmp_path / "saida")
    app.iniciar_captura()
    assert not ativo(app.btn_iniciar) and ativo(app.btn_cancelar)
    esperar_fim(app)

    assert [r.id_lattes for r in app.registros] == [ID_A, ID_B, ID_C]
    assert avisos[-1][:2] == ("showinfo", "Captura concluída")
    assert app.label_progresso.cget("text") == "3 / 3"
    assert ativo(app.btn_exportar)

    arquivos = set(os.listdir(app.pasta_execucao))
    assert {coletor.NOME_PLANILHA, coletor.NOME_CSV, coletor.NOME_LOG, f"{ID_A}.png"} <= arquivos
    ws = openpyxl.load_workbook(os.path.join(app.pasta_execucao, coletor.NOME_PLANILHA)).active
    assert [ws.cell(row=i, column=2).value for i in (2, 3, 4)] == [ID_A, ID_B, ID_C]


def test_captura_com_problemas_avisa_quais_ids(app, tmp_path, avisos, coletor_falso):
    coletor_falso.roteiro = {ID_B: coletor.STATUS_CAPTCHA}
    preparar(app, tmp_path / "saida")
    app.iniciar_captura()
    esperar_fim(app)
    tipo, titulo, texto = avisos[-1]
    assert (tipo, titulo) == ("showwarning", "Captura concluída com problemas")
    assert f"{ID_B}: {coletor.STATUS_CAPTCHA}" in texto


def test_cancelar_e_continuar(app, tmp_path, avisos, coletor_falso):
    coletor_falso.bloquear = True
    preparar(app, tmp_path / "saida")
    app.iniciar_captura()
    app.janela.update()
    app.cancelar_captura()
    esperar_fim(app)

    assert avisos[-1][:2] == ("showinfo", "Captura cancelada")
    assert app.btn_iniciar.cget("text") == "▶  Continuar (3 restantes)"

    coletor_falso.bloquear = False
    pasta_anterior = app.pasta_execucao
    app.iniciar_captura()
    esperar_fim(app)
    assert app.pasta_execucao == pasta_anterior  # retoma na mesma subpasta
    assert len(app.registros) == 3


def test_pausa_depois_de_falhas_seguidas(app, tmp_path, avisos, coletor_falso):
    coletor_falso.roteiro = dict.fromkeys((ID_A, ID_B, ID_C), coletor.STATUS_ERRO)
    preparar(app, tmp_path / "saida", ids=(ID_A, ID_B, ID_C, "9999888877776666"))
    app.iniciar_captura()
    esperar_fim(app)
    assert avisos[-1][:2] == ("showwarning", "Captura pausada")
    assert app.btn_iniciar.cget("text") == "▶  Continuar (1 restantes)"


# --- erros inesperados ------------------------------------------------------------

@pytest.fixture
def log_de_erros(monkeypatch, tmp_path):
    caminho = tmp_path / "config" / "lattesink.log"
    monkeypatch.setattr(LattesINK, "ARQUIVO_LOG_ERROS", str(caminho))
    handler = LattesINK.configurar_log_de_erros()
    yield caminho
    LattesINK.logger.removeHandler(handler)
    handler.close()


def test_erro_num_botao_vira_aviso_e_vai_para_o_log(app, avisos, log_de_erros):
    def botao_com_defeito():
        raise ZeroDivisionError("divisão por zero")

    app.janela.after(0, botao_com_defeito)
    app.janela.update()

    tipo, titulo, texto = avisos[-1]
    assert (tipo, titulo) == ("showerror", "Erro inesperado")
    assert "divisão por zero" in texto and str(log_de_erros) in texto
    conteudo = log_de_erros.read_text(encoding="utf-8")
    assert "Traceback" in conteudo and "ZeroDivisionError" in conteudo
    assert f"LattesINK {LattesINK.__version__}" in conteudo


def test_log_de_erros_so_e_criado_quando_preciso(log_de_erros):
    assert not log_de_erros.exists()


# --- abrir a pasta ----------------------------------------------------------------

def test_abrir_pasta(app, tmp_path, monkeypatch, avisos):
    abertas = []
    monkeypatch.setattr(LattesINK, "abrir_no_sistema", abertas.append)
    assert not ativo(app.btn_abrir)

    preparar(app, tmp_path / "saida")
    assert ativo(app.btn_abrir)
    app.abrir_pasta()
    assert abertas == [str(tmp_path / "saida")]  # antes da captura: a pasta de destino

    app.iniciar_captura()
    esperar_fim(app)
    app.abrir_pasta()
    assert abertas[-1] == app.pasta_execucao  # depois: a subpasta da captura

    def falha(_):
        raise OSError("sem gerenciador de arquivos")

    monkeypatch.setattr(LattesINK, "abrir_no_sistema", falha)
    app.abrir_pasta()
    assert avisos[-1][:2] == ("showerror", "Erro ao abrir a pasta")


# --- repetir as falhas --------------------------------------------------------------

def test_repetir_so_os_que_falharam(app, tmp_path, avisos, coletor_falso):
    coletor_falso.roteiro = {ID_B: [coletor.STATUS_CAPTCHA, coletor.STATUS_OK]}
    preparar(app, tmp_path / "saida")
    app.iniciar_captura()
    esperar_fim(app)
    assert ativo(app.btn_repetir)
    assert app.btn_repetir.cget("text") == "Repetir falhas (1)"
    pasta = app.pasta_execucao

    app.repetir_falhas()
    esperar_fim(app)

    assert app.pasta_execucao == pasta
    assert sorted(r.id_lattes for r in app.registros) == sorted((ID_A, ID_B, ID_C))
    assert all(r.status == coletor.STATUS_OK for r in app.registros)
    assert not ativo(app.btn_repetir)
    assert avisos[-1][:2] == ("showinfo", "Captura concluída")

    # a planilha sai na ordem do arquivo de IDs, mesmo com o ID_B capturado por último
    ws = openpyxl.load_workbook(os.path.join(pasta, coletor.NOME_PLANILHA)).active
    assert [ws.cell(row=i, column=2).value for i in (2, 3, 4)] == [ID_A, ID_B, ID_C]
    # o CSV de backup guarda as duas tentativas
    with open(os.path.join(pasta, coletor.NOME_CSV), encoding="utf-8-sig") as f:
        assert sum(ID_B in linha for linha in f) == 2


def test_repetir_falhas_fica_desativado_sem_falhas(app, tmp_path):
    preparar(app, tmp_path / "saida")
    assert not ativo(app.btn_repetir)
    app.iniciar_captura()
    esperar_fim(app)
    assert not ativo(app.btn_repetir)
    assert app.btn_repetir.cget("text") == "Repetir falhas"
